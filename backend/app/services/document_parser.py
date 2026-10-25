from pathlib import Path
from docx import Document as DocxDocument
import fitz
from PIL import Image
import pytesseract
from io import BytesIO

def parse_document(path:Path, suffix:str, ocr_enabled:bool=False):
    suffix=suffix.lower()
    if suffix==".pdf":
        doc=fitz.open(path); pages=[(i+1,p.get_text("text")) for i,p in enumerate(doc)]; doc.close()
        if not any(t.strip() for _,t in pages): raise ValueError("No selectable text found. This may be a scanned PDF; configure OCR to process it.")
        return pages
    if suffix==".docx":
        doc=DocxDocument(path); texts=[p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows: texts.append(" | ".join(c.text.strip() for c in row.cells))
        if not texts: raise ValueError("The DOCX file contains no readable text.")
        return [(None,"\n".join(texts))]
    if suffix==".txt":
        raw=path.read_bytes()
        for enc in ("utf-8-sig","utf-16","cp1252"):
            try: value=raw.decode(enc); break
            except UnicodeDecodeError: continue
        else: raise ValueError("Could not decode text file.")
        if not value.strip(): raise ValueError("The text file is empty.")
        return [(None,value)]
    if suffix in {".png",".jpg",".jpeg",".webp"}:
        if not ocr_enabled: raise ValueError("Image stored, but OCR is disabled. Install Tesseract and set OCR_ENABLED=true to extract image text; visual understanding requires a vision provider.")
        try: value=pytesseract.image_to_string(Image.open(path))
        except Exception as exc: raise ValueError("OCR failed. Check that Tesseract is installed and configured.") from exc
        if not value.strip(): raise ValueError("OCR found no readable text in this image.")
        return [(None,value)]
    raise ValueError("Unsupported file type.")
